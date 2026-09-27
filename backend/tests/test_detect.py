import io

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app.detection import analyse
from app.main import app

ENGLISH = ("Government Resolution on crop insurance. Farmers with land under two hectares "
           "are eligible for the scheme from the next kharif season. ") * 4
HINDI = "किसानों के लिए फसल बीमा योजना सर्वे क्रमांक १४२ पर लागू है। " * 6


def pdf_with_text() -> bytes:
    doc = pymupdf.open()
    for _ in range(2):
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), ENGLISH, fontsize=11)
    return doc.tobytes()


def pdf_scanned() -> bytes:
    """A PDF whose only page is an image, like a scanner's output."""
    img = io.BytesIO()
    Image.new("L", (600, 800), 255).save(img, format="PNG")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=img.getvalue())
    return doc.tobytes()


def docx_bytes() -> bytes:
    d = docx.Document()
    d.add_heading("Circular on ATKT rules", 1)
    d.add_paragraph(ENGLISH)
    d.add_paragraph(HINDI)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (320, 200), (40, 40, 40)).save(buf, format="PNG", dpi=(300, 300))
    return buf.getvalue()


MARKDOWN = (b"# Sakshya notes\n\n## Scope\n\nWe ingest **PDF**, *DOCX* and scans.\n\n"
            b"- English\n- Hindi\n- Marathi\n\n```python\nprint('hi')\n```\n\n"
            b"See [the plan](docs/Review2_Plan.md) for details.\n")


def test_native_pdf():
    r = analyse(pdf_with_text(), "gr.pdf")
    assert r["type"]["label"] == "pdf"
    assert r["type"]["kind"] == "PDF"
    assert r["pdf_text_layer"] == "native"
    assert r["ingestion"]["status"] == "in_scope"
    assert r["extension_check"]["status"] == "match"
    assert {"label": "Pages", "value": "2"} in r["facts"]


def test_scanned_pdf_goes_to_ocr():
    r = analyse(pdf_scanned(), "scan.pdf")
    assert r["type"]["label"] == "pdf"
    assert r["pdf_text_layer"] == "scanned"
    assert "OCR" in r["ingestion"]["route"]


def test_docx_with_devanagari():
    r = analyse(docx_bytes(), "circular.docx")
    assert r["type"]["label"] == "docx"
    assert r["type"]["kind"] == "Word document"
    assert r["scripts"]["devanagari"] > 0.2
    assert r["scripts"]["devanagari_digits"] is True


def test_png_image():
    r = analyse(png_bytes(), "page.png")
    assert r["type"]["kind"] == "Image"
    assert {"label": "Dimensions", "value": "320 × 200 px"} in r["facts"]
    assert r["ingestion"]["status"] == "in_scope"


def test_markdown():
    r = analyse(MARKDOWN, "notes.md")
    assert r["type"]["label"] == "markdown"
    assert r["type"]["stamp"] == "MD"
    assert r["ingestion"]["status"] == "out_of_scope"


def test_hindi_plain_text():
    r = analyse(HINDI.encode("utf-8"), "note.txt")
    assert r["type"]["is_text"] is True
    assert r["scripts"]["devanagari"] > 0.9


def test_content_wins_over_a_wrong_extension():
    r = analyse(docx_bytes(), "looks-like.pdf")
    assert r["type"]["label"] == "docx"
    assert r["extension_check"]["status"] == "mismatch"
    assert r["extension_check"]["expected_kind"] == "PDF"


def test_empty_file():
    r = analyse(b"", "blank.pdf")
    assert r["type"]["label"] == "empty"
    assert r["detection"]["method"] == "size"


def test_tiny_pdf_falls_back_to_signature():
    r = analyse(b"%PDF-1.4\n%%EOF", "tiny.pdf")
    assert r["type"]["label"] == "pdf"
    assert r["detection"]["method"] in {"content", "signature"}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def test_health(client):
    body = client.get("/api/health").json()
    assert body["status"] == "ok"
    assert body["detector"]["name"] == "Magika"


def test_detect_endpoint_multiple_files(client):
    files = [
        ("files", ("gr.pdf", pdf_with_text(), "application/pdf")),
        ("files", ("notes.md", MARKDOWN, "text/markdown")),
        ("files", ("page.png", png_bytes(), "image/png")),
    ]
    results = client.post("/api/detect", files=files).json()["results"]
    assert [r["type"]["label"] for r in results] == ["pdf", "markdown", "png"]
