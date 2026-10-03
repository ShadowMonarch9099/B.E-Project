"""Devanagari typed in legacy (non-Unicode) fonts must not be trusted as a text layer."""

import io

import docx
import pymupdf
import pytest

from app.detection import analyse
from app.legacy_fonts import is_legacy_font, legacy_spans, looks_legacy
from app.ocr import engine_status, run_ocr
from test_ocr import png_bytes, scan_image

# "भारत का संविधान" as stored by the DVB-TT Surekh font: the title of the official Hindi
# Constitution PDF (legislative.gov.in, 2024) as it appears in its text layer.
DVB = "£ÉÉ®iÉ BÉEÉ ºÉÆÉÊ´ÉvÉÉxÉ [1 , 2024 ] THE CONSTITUTION OF INDIA"
# "किसानों के लिए फसल बीमा योजना अगले खरीफ मौसम से लागू होगी। ..." in Kruti Dev.
KRUTI = ("fdlkuksa ds fy, Qly chek ;kstuk vxys [kjhQ ekSle ls ykxw gksxhA ljdkj us dgk fd ;g "
         ";kstuk Hkkjr esa lHkh fdlkuksa ds fy, gS vkSj bldk ykHk gj fdlku dks feysxkA")


@pytest.mark.parametrize("text", [DVB * 3, KRUTI])
def test_legacy_text_is_recognised(text):
    assert looks_legacy(text)


@pytest.mark.parametrize("text", [
    "The Government of Maharashtra has issued a resolution on crop insurance for farmers. " * 3,
    "Die Bundesregierung hat beschlossen, dass die Maßnahmen für Übergänge verlängert werden.",
    "Le gouvernement a décidé que les élèves bénéficieront d'une réduction après les vacances d'été.",
    "if dk > ds: print(dh); x = ls -la | grep esa; fd = open(ij)",
    "किसानों के लिए फसल बीमा योजना अगले खरीफ मौसम से लागू होगी। सर्वे क्रमांक १४२",
])
def test_normal_text_is_not_legacy(text):
    assert not looks_legacy(text)


def test_font_names():
    assert is_legacy_font("ABCDEF+DVBW-TTSurekhNormal")
    assert is_legacy_font("KrutiDev010") and is_legacy_font("Shivaji01")
    assert not is_legacy_font("Mangal") and not is_legacy_font("NirmalaUI")
    assert not is_legacy_font("Arial-BoldMT") and not is_legacy_font("Lohit-Devanagari")


class FakePage:
    def __init__(self, spans):
        self.spans = spans

    def get_text(self, kind):
        return {"blocks": [{"lines": [{"spans": [{"font": f, "text": t} for f, t in self.spans]}]}]}


def test_legacy_spans_counts_only_non_unicode_text():
    page = FakePage([
        ("ABCDEF+DVBW-TTSurekhNormal", "£ÉÉ®iÉ BÉEÉ ºÉÆÉÊ´ÉvÉÉxÉ"),
        ("Shivaji05", "भारत का संविधान"),  # legacy name but Unicode text: fine
        ("Arial", "THE CONSTITUTION"),
    ])
    legacy, total, fonts = legacy_spans(page)
    assert fonts == {"DVBW-TTSurekhNormal"}
    assert legacy == len("£ÉÉ®iÉ BÉEÉ ºÉÆÉÊ´ÉvÉÉxÉ") and total > legacy


def legacy_pdf(hidden_text: str, with_english_page: bool = False) -> bytes:
    """A page that looks like Devanagari but whose text layer is legacy-font nonsense."""
    doc = pymupdf.open()
    if with_english_page:
        doc.new_page().insert_textbox(pymupdf.Rect(50, 50, 550, 800),
                                      "Government of Maharashtra, Revenue Department. " * 12)
    page = doc.new_page()
    page.insert_image(page.rect, stream=png_bytes(scan_image()))
    page.insert_textbox(pymupdf.Rect(60, 60, 540, 780), hidden_text, fontname="helv", render_mode=3)
    return doc.tobytes()


@pytest.mark.parametrize("hidden", [DVB * 4, KRUTI])
def test_legacy_pdf_is_sent_to_ocr(hidden):
    r = analyse(legacy_pdf(hidden), "gr.pdf")
    assert r["pdf_text_layer"] == "legacy"
    assert r["ingestion"]["needs_ocr"] is True
    assert "legacy" in r["ingestion"]["route"]
    assert r["scripts"] is None  # the nonsense text must not count as Latin script
    assert any(f["label"] == "Pages in a legacy Devanagari font" for f in r["facts"])


def test_bilingual_pdf_with_one_legacy_page_is_mixed():
    r = analyse(legacy_pdf(KRUTI, with_english_page=True), "gr.pdf")
    assert r["pdf_text_layer"] == "mixed"
    assert r["ingestion"]["needs_ocr"] is True


@pytest.mark.skipif(not engine_status()["available"], reason="Tesseract with eng/hin/mar is not installed")
def test_ocr_reads_legacy_pages_from_the_image():
    r = run_ocr(legacy_pdf(KRUTI, with_english_page=True), "pdf")
    first, second = r["pages"]
    assert first["source"] == "text_layer"
    assert second["source"] == "ocr" and second["ocr_reason"] == "legacy_font"
    assert "महाराष्ट्र" in second["text"]
    assert r["summary"]["legacy_font_pages"] == [2]


def test_word_file_in_kruti_dev_is_flagged():
    d = docx.Document()
    run = d.add_paragraph().add_run(KRUTI)
    run.font.name = "Kruti Dev 010"
    buf = io.BytesIO()
    d.save(buf)
    r = analyse(buf.getvalue(), "circular.docx")
    assert {"label": "Legacy Devanagari font", "value": "Kruti Dev 010"} in r["facts"]
    assert "legacy" in r["note"] and r["scripts"] is None


# The first lines of a sample GR's text layer: Unicode Devanagari, but the conjuncts are
# stray Latin Extended / IPA characters ("राष्ट्रीय" -> "राʊȄीय", "प्रकल्प" -> "ɓकȧप").
BROKEN = ("राʊȄीय उɧतर शʺȈा अभʺयान (RUSA) या योजनेअंतȁगत राȑय ɓकȧप संचालनालयास निधी उपलȡध "
          "कʕन देȘयाबाबत. महाराʊȄ शासन उɧ व तंɑ शʺȈण विभाग शासन निȁणय Ɍमांक: ʖसायो २०२३")


def test_broken_conjuncts_are_recognised():
    from app.legacy_fonts import looks_broken

    assert looks_broken(BROKEN)
    assert not looks_legacy(BROKEN)  # a different failure from a legacy font


@pytest.mark.parametrize("text", [
    "महाराष्ट्र शासन उच्च व तंत्र शिक्षण विभाग राष्ट्रीय उच्चतर शिक्षा अभियान प्रकल्प संचालनालय",
    "Phonetics: the vowel ə and the glide ʊ appear in English words such as about and foot. " * 2,
    "भारत का संविधान THE CONSTITUTION OF INDIA ə ʊ",  # too few Devanagari words to judge
])
def test_clean_text_is_not_broken(text):
    from app.legacy_fonts import looks_broken

    assert not looks_broken(text)


def test_pdf_with_broken_devanagari_goes_to_ocr():
    from tests.helpers import html_pdf

    # PyMuPDF's HTML writer with Lohit Devanagari produces exactly this kind of text layer.
    text = "<p>महाराष्ट्र शासन उच्च व तंत्र शिक्षण विभाग. राष्ट्रीय उच्चतर शिक्षा अभियान प्रकल्प संचालनालय.</p>" * 6
    r = analyse(html_pdf([text]), "gr.pdf")
    assert r["pdf_text_layer"] == "broken" and r["ingestion"]["needs_ocr"] is True
    assert {"label": "Pages with broken Devanagari text", "value": "1 of 1"} in r["facts"]
