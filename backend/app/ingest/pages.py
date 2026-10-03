"""Read every page of a document into text with layout (plan section 6.2, steps 1-3).

Pages with a usable text layer are read with PyMuPDF. Pages with no text layer, or with
Devanagari in a legacy (non-Unicode) font, are rendered at 300 DPI and OCR'd with Tesseract.
Every word keeps its box in PDF points (top-left origin, the same space the PDF viewer
draws in), so a chunk's evidence can be highlighted on the page later.

The layout is paragraphs -> lines -> words; each word is {"t": text, "b": box or None,
"c": OCR confidence 0-100 or None}. `build_text` joins it into the page text and records
each word's start offset ("s"), which ties sentences and chunks back to their boxes.
"""

import io
from collections.abc import Callable, Iterator
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from dataclasses import dataclass, field

from PIL import Image

from .. import config
from ..inspect_file import classify_page
from ..ocr import OCR_REASON, PDF_DPI, ocr_image


@dataclass
class ExtractedPage:
    page_no: int
    source: str  # native / ocr / docx
    width: float | None
    height: float | None
    paragraphs: list[list[list[dict]]] = field(default_factory=list)
    ocr_conf: float | None = None  # 0-1
    ocr_reason: str | None = None  # no_text_layer / legacy_font / broken_text
    check_reason: str | None = None  # low_confidence / no_text / ocr_unavailable

    def build(self) -> tuple[str, list]:
        return build_text(self.paragraphs)


def build_text(paragraphs: list[list[list[dict]]]) -> tuple[str, list]:
    """Join words with spaces, lines with newlines and paragraphs with blank lines."""
    parts: list[str] = []
    pos = 0
    layout: list = []
    for para in paragraphs:
        lines = [line for line in para if line]
        if not lines:
            continue
        if layout:
            parts.append("\n\n")
            pos += 2
        out_para = []
        for li, line in enumerate(lines):
            if li:
                parts.append("\n")
                pos += 1
            out_line = []
            for wi, word in enumerate(line):
                if wi:
                    parts.append(" ")
                    pos += 1
                out_line.append({**word, "s": pos})
                parts.append(word["t"])
                pos += len(word["t"])
            out_para.append(out_line)
        layout.append(out_para)
    return "".join(parts), layout


def _round_box(x0, y0, x1, y1) -> list[float]:
    return [round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1)]


def native_page(page, page_no: int, source: str = "native") -> ExtractedPage:
    """Words from the text layer, grouped by PyMuPDF's blocks (paragraphs) and lines."""
    import pymupdf

    rotate = page.rotation_matrix if page.rotation else None
    paragraphs: list[list[list[dict]]] = []
    last_block = last_line = None
    for x0, y0, x1, y1, text, block, line, _ in page.get_text("words"):
        if not text.strip():
            continue
        if rotate is not None:  # text coordinates are unrotated; the viewer shows the rotated page
            r = pymupdf.Rect(x0, y0, x1, y1) * rotate
            x0, y0, x1, y1 = r.x0, r.y0, r.x1, r.y1
        word = {"t": text, "b": _round_box(x0, y0, x1, y1), "c": None}
        if block != last_block:
            paragraphs.append([[word]])
        elif line != last_line:
            paragraphs[-1].append([word])
        else:
            paragraphs[-1][-1].append(word)
        last_block, last_line = block, line
    rect = page.rect
    return ExtractedPage(page_no, source, round(rect.width, 2), round(rect.height, 2), paragraphs)


def _ocr_page(image: Image.Image, page_no: int, width: float, height: float, reason: str) -> ExtractedPage:
    result = ocr_image(image, page_no)
    sx = width / result.width
    sy = height / result.height
    paragraphs: list[list[list[dict]]] = []
    last_par = last_line = None
    for w in result.layout_words:
        block, par, line = w["key"]
        x0, y0, x1, y1 = w["bbox"]
        word = {"t": w["text"], "b": _round_box(x0 * sx, y0 * sy, x1 * sx, y1 * sy), "c": w["conf"]}
        if (block, par) != last_par:
            paragraphs.append([[word]])
        elif line != last_line:
            paragraphs[-1].append([word])
        else:
            paragraphs[-1][-1].append(word)
        last_par, last_line = (block, par), line
    conf = None if result.confidence is None else round(result.confidence / 100, 4)
    return ExtractedPage(page_no, "ocr", width, height, paragraphs, conf, reason, result.check_reason())


def extract_pdf(path, ocr_available: bool, source_for_text: str = "native",
                on_total: Callable[[int], None] | None = None) -> Iterator[ExtractedPage]:
    """Yield every page of a PDF, in completion order (OCR'd pages run in parallel)."""
    import pymupdf

    pymupdf.TOOLS.mupdf_display_errors(False)
    workers = max(1, config.OCR_WORKERS)
    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise ValueError("The PDF is password-protected.")
        if on_total:
            on_total(doc.page_count)
        pending: dict = {}
        with ThreadPoolExecutor(max_workers=workers) as pool:
            for i in range(doc.page_count):
                page = doc[i]
                kind, _, _ = classify_page(page)
                if kind == "text":
                    yield native_page(page, i + 1, source_for_text)
                    continue
                rect = page.rect
                reason = OCR_REASON[kind]
                if not ocr_available:
                    yield ExtractedPage(i + 1, "ocr", round(rect.width, 2), round(rect.height, 2),
                                        ocr_reason=reason, check_reason="ocr_unavailable")
                    continue
                # Render here: a PyMuPDF document must not be shared between threads.
                pix = page.get_pixmap(dpi=PDF_DPI, colorspace=pymupdf.csGRAY)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                fut = pool.submit(_ocr_page, image, i + 1, round(rect.width, 2), round(rect.height, 2), reason)
                pending[fut] = i + 1
                while len(pending) >= workers * 2:  # bound the number of rendered pages in memory
                    done, _ = wait(pending, return_when=FIRST_COMPLETED)
                    for f in done:
                        pending.pop(f)
                        yield f.result()
            while pending:
                done, _ = wait(pending, return_when=FIRST_COMPLETED)
                for f in done:
                    pending.pop(f)
                    yield f.result()


def docx_paragraphs(path) -> ExtractedPage:
    """A DOCX read paragraph by paragraph (no pages), used when it cannot be converted to PDF."""
    import docx

    document = docx.Document(str(path))
    texts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                texts.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat their text
    paragraphs = [[[{"t": w, "b": None, "c": None} for w in line.split()]
                   for line in text.splitlines() if line.strip()] for text in texts]
    return ExtractedPage(1, "docx", None, None, paragraphs)


def image_to_pdf(data: bytes, ext: str) -> bytes:
    """Wrap an image (or a multi-page TIFF) in a PDF so it opens in the PDF viewer."""
    import pymupdf

    with pymupdf.open(stream=data, filetype=ext) as img:
        return img.convert_to_pdf()
