"""Non-ML facts about a file once its type is known.

These are plain parsers (PyMuPDF, python-docx, Pillow, zipfile), not models.
The PDF text-layer check follows plan section 6.1 step 1: a page counts as
having a usable text layer when it has enough characters, most of them are
valid Latin or Devanagari text, and it is not Devanagari typed in a legacy
(non-Unicode) font. Pages without a usable text layer go to OCR.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field

from .legacy_fonts import is_legacy_font, legacy_spans, looks_broken, looks_legacy

MAX_PDF_PAGES_CHECKED = 60
MAX_TEXT_CHARS = 400_000
MIN_PAGE_CHARS = 20
MIN_VALID_SHARE = 0.5
LEGACY_SPAN_SHARE = 0.2  # a page is legacy when this much of its text is in legacy fonts


@dataclass
class Facts:
    items: list[tuple[str, str]] = field(default_factory=list)
    scripts: dict | None = None
    pdf_text_layer: str | None = None
    encrypted: bool = False
    note: str | None = None

    def add(self, label: str, value) -> None:
        if value not in (None, ""):
            self.items.append((label, str(value)))


def _is_devanagari(ch: str) -> bool:
    return "ऀ" <= ch <= "ॿ" or "꣠" <= ch <= "ꣿ"


def _is_latin_letter(ch: str) -> bool:
    return ("a" <= ch.lower() <= "z") or ("À" <= ch <= "ɏ" and ch.isalpha())


def script_mix(text: str) -> dict | None:
    """Share of Devanagari, Latin and other letters, plus whether Devanagari digits appear."""
    text = text[:MAX_TEXT_CHARS]
    deva = latin = other = 0
    deva_digits = False
    for ch in text:
        if "०" <= ch <= "९":
            deva_digits = True
        elif _is_devanagari(ch):
            deva += 1
        elif _is_latin_letter(ch):
            latin += 1
        elif ch.isalpha():
            other += 1
    total = deva + latin + other
    if total == 0:
        return None
    return {
        "devanagari": round(deva / total, 3),
        "latin": round(latin / total, 3),
        "other": round(other / total, 3),
        "devanagari_digits": deva_digits,
    }


def _valid_share(text: str) -> float:
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return 0.0
    valid = sum(1 for c in chars if _is_devanagari(c) or _is_latin_letter(c) or c.isdigit()
                or c in ".,;:!?'\"()[]-/%&@#*+=_₹")
    return valid / len(chars)


def page_has_text_layer(text: str) -> bool:
    stripped = text.strip()
    return len(stripped) >= MIN_PAGE_CHARS and _valid_share(stripped) >= MIN_VALID_SHARE


def classify_page(page) -> tuple[str, str, set[str]]:
    """(kind, text, legacy font names) for a PyMuPDF page; kind is "text", "legacy", "broken" or "none".

    "legacy" means the page has a text layer, but it is Devanagari in a pre-Unicode font.
    "broken" means it is Unicode Devanagari with conjuncts stored as stray characters
    ("राष्ट्रीय" as "राʊȄीय"). Either way the text cannot be used and the page must be OCR'd.
    """
    text = page.get_text()
    if not page_has_text_layer(text):
        return "none", text, set()
    legacy_chars, total, fonts = legacy_spans(page)
    if (total and legacy_chars / total >= LEGACY_SPAN_SHARE) or looks_legacy(text):
        return "legacy", text, fonts
    if looks_broken(text):
        return "broken", text, set()
    return "text", text, set()


def _pdf(data: bytes, facts: Facts) -> None:
    import pymupdf

    pymupdf.TOOLS.mupdf_display_errors(False)  # malformed-but-readable PDFs otherwise flood the log
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            facts.encrypted = True
            facts.add("Pages", doc.page_count)
            facts.add("Encryption", "Password protected")
            return
        pages = doc.page_count
        facts.add("Pages", pages)
        checked = min(pages, MAX_PDF_PAGES_CHECKED)
        counts = {"text": 0, "legacy": 0, "broken": 0, "none": 0}
        legacy_fonts: set[str] = set()
        texts: list[str] = []  # only usable (Unicode) text feeds the script mix
        for i in range(checked):
            kind, text, fonts = classify_page(doc[i])
            counts[kind] += 1
            legacy_fonts |= fonts
            if kind == "text":
                texts.append(text)
        if checked:
            facts.pdf_text_layer = ("native" if counts["text"] == checked
                                    else "scanned" if counts["none"] == checked
                                    else "legacy" if counts["legacy"] == checked
                                    else "broken" if counts["broken"] == checked else "mixed")
            suffix = "" if checked == pages else f" (first {checked} checked)"
            facts.add("Pages with a usable text layer", f"{counts['text']} of {checked}{suffix}")
            if counts["legacy"]:
                facts.add("Pages in a legacy Devanagari font", f"{counts['legacy']} of {checked}{suffix}")
                facts.add("Legacy fonts", ", ".join(sorted(legacy_fonts)) or "not named; found from the text")
            if counts["broken"]:
                facts.add("Pages with broken Devanagari text", f"{counts['broken']} of {checked}{suffix}")
        meta = doc.metadata or {}
        facts.add("Title", (meta.get("title") or "").strip())
        facts.add("Author", (meta.get("author") or "").strip())
        facts.add("Created with", (meta.get("producer") or meta.get("creator") or "").strip())
        facts.scripts = script_mix("".join(texts))


def _docx(data: bytes, facts: Facts) -> None:
    import docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    table_text = [c.text for t in document.tables for r in t.rows for c in r.cells]
    text = "\n".join(paragraphs + table_text)
    facts.add("Paragraphs", len(paragraphs))
    facts.add("Tables", len(document.tables))
    facts.add("Words", len(text.split()))
    facts.add("Title", (document.core_properties.title or "").strip())
    facts.add("Author", (document.core_properties.author or "").strip())

    runs = [r for p in document.paragraphs for r in p.runs if r.text.strip()]
    fonts = {r.font.name for r in runs if r.font.name and is_legacy_font(r.font.name)}
    legacy_chars = sum(len(r.text) for r in runs if r.font.name in fonts)
    total_chars = sum(len(r.text) for r in runs)
    if (total_chars and legacy_chars / total_chars >= LEGACY_SPAN_SHARE) or looks_legacy(text):
        facts.add("Legacy Devanagari font", ", ".join(sorted(fonts)) or "found from the text")
        facts.note = ("This Word file is typed in a legacy (non-Unicode) Devanagari font, so its text "
                      "is not real Hindi or Marathi. Save it as PDF and upload that to OCR it.")
        return
    facts.scripts = script_mix(text)


def _office_zip(label: str, data: bytes, facts: Facts) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
    if label == "xlsx":
        facts.add("Sheets", sum(1 for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)))
    elif label == "pptx":
        facts.add("Slides", sum(1 for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)))
    elif label in {"zip", "epub"}:
        facts.add("Entries", len(names))


def _image(data: bytes, facts: Facts) -> None:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as im:
        facts.add("Dimensions", f"{im.width} × {im.height} px")
        facts.add("Colour mode", im.mode)
        dpi = im.info.get("dpi")
        if dpi:
            facts.add("Resolution", f"{round(float(dpi[0]))} DPI")
        frames = getattr(im, "n_frames", 1)
        if frames > 1:
            facts.add("Frames / pages", frames)


def decode_text(data: bytes) -> tuple[str, str]:
    """Decode text as UTF-16 (with BOM) or UTF-8, else fall back to Windows-1252."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace"), "UTF-16"
    try:
        return data.decode("utf-8-sig"), "UTF-8"
    except UnicodeDecodeError as exc:
        # A multi-byte character cut off at the end of a truncated sample is still UTF-8.
        if exc.start >= len(data) - 3:
            return data[: exc.start].decode("utf-8-sig"), "UTF-8"
        return data.decode("cp1252", errors="replace"), "Windows-1252 (guessed)"


def _text(label: str, data: bytes, facts: Facts) -> None:
    text, encoding = decode_text(data[: MAX_TEXT_CHARS * 4])
    facts.add("Encoding", encoding)
    facts.add("Lines", text.count("\n") + (1 if text and not text.endswith("\n") else 0))
    facts.add("Words", len(text.split()))
    if label == "markdown":
        facts.add("Headings", len(re.findall(r"^#{1,6}\s", text, flags=re.M)))
    facts.scripts = script_mix(text)


def inspect(label: str, is_text: bool, group: str, data: bytes) -> Facts:
    facts = Facts()
    try:
        if label == "pdf":
            _pdf(data, facts)
        elif label == "docx":
            _docx(data, facts)
        elif label in {"xlsx", "pptx", "zip", "epub"}:
            _office_zip(label, data, facts)
        elif group == "image" and label != "svg":
            _image(data, facts)
        elif is_text:
            _text(label, data, facts)
    except Exception as exc:  # a damaged or unusual file should not fail the whole upload
        facts.note = f"Type detected, but the file could not be read further ({type(exc).__name__})."
    return facts
